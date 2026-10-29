import io

import pytest

from core.documents import read_document, read_path


def test_reads_plain_text():
    text = "Amina Bello\nData analyst with three years of experience."
    assert read_document("cv.txt", text.encode("utf-8")) == text


def test_reads_docx_paragraphs_and_tables():
    import docx

    doc = docx.Document()
    doc.add_paragraph("Amina Bello, data analyst in Lagos")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Excel and SQL"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = read_document("cv.docx", buffer.getvalue())
    assert "Amina Bello" in text
    assert "Skill | Excel and SQL" in text


def test_unsupported_extension_gives_a_clear_error():
    with pytest.raises(ValueError, match="Unsupported file type"):
        read_document("cv.doc", b"anything")


def test_nearly_empty_file_gives_a_clear_error():
    with pytest.raises(ValueError, match="almost no readable text"):
        read_document("cv.txt", b"hi")


def test_read_path_reads_from_disk(tmp_path):
    path = tmp_path / "cv.txt"
    path.write_text("Chioma Eze, monitoring and evaluation assistant in Enugu.", encoding="utf-8")
    assert "Chioma Eze" in read_path(path)
