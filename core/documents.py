"""Read text out of uploaded documents (PDF, DOCX, TXT)."""

import io
from pathlib import Path


def read_document(name, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    elif ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for tbl in d.tables:
            for row in tbl.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in (".txt", ".md", ""):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT (save .doc files as .docx).")
    if len(text.strip()) < 30:
        raise ValueError(f"'{name}' has almost no readable text. If it is a scanned PDF, run OCR first.")
    return text


def read_path(path):
    path = Path(path)
    return read_document(path.name, path.read_bytes())
